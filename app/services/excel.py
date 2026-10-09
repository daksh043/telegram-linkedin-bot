"""Excel tracker export via openpyxl.

SQLite remains the source of truth; this workbook is the human-friendly
tracker described in EXCEL_TRACKER_SPEC.md.
"""

from __future__ import annotations

import asyncio
import os
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.config import Config
from app.database.repository import Repository
from app.services.analytics import compute_analytics
from app.utils.dates import fmt, now_ts

TRACKER_HEADERS = [
    "ID", "Company", "Contact Person", "Role", "LinkedIn", "Email",
    "Industry", "Opportunity", "Personalization Angle", "Message",
    "Channel", "Status", "Date Added", "Date Contacted",
    "Follow-up Date", "Response", "Notes",
]

CONTACT_HEADERS = [
    "Contact ID", "Company ID", "Company", "Name", "Role",
    "LinkedIn", "Email", "Notes",
]

DASHBOARD_ROWS = [
    "Total Companies", "Total Contacts", "Drafted", "Sent", "Replied",
    "Interested", "Meetings", "Follow-ups", "Not Interested", "Closed",
    "Response Rate", "Interest Rate",
]

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


@dataclass
class ExportData:
    tracker_rows: list[list] = field(default_factory=list)
    contact_rows: list[list] = field(default_factory=list)
    dashboard_values: dict[str, object] = field(default_factory=dict)


def build_export_data(
    companies: list[dict],
    contacts: list[dict],
    outreach: list[dict],
    analytics: dict,
) -> ExportData:
    companies_by_id = {c["id"]: c for c in companies}
    contacts_by_id = {c["id"]: c for c in contacts}

    tracker_rows = []
    for o in outreach:
        company = companies_by_id.get(o["company_id"], {})
        contact = contacts_by_id.get(o["contact_id"], {}) if o["contact_id"] else {}
        tracker_rows.append([
            o["id"],
            company.get("name", ""),
            contact.get("name", ""),
            contact.get("role", ""),
            contact.get("linkedin_url", ""),
            contact.get("email", ""),
            company.get("industry", ""),
            company.get("opportunity", ""),
            company.get("personalization_angle", ""),
            o.get("message", ""),
            o.get("channel", ""),
            o.get("status", ""),
            fmt(o.get("date_added")),
            fmt(o.get("date_contacted")),
            fmt(o.get("follow_up_date")),
            o.get("response", ""),
            o.get("notes", ""),
        ])

    contact_rows = [
        [
            c["id"],
            c["company_id"],
            companies_by_id.get(c["company_id"], {}).get("name", ""),
            c["name"],
            c.get("role", ""),
            c.get("linkedin_url", ""),
            c.get("email", ""),
            c.get("notes", ""),
        ]
        for c in contacts
    ]

    dashboard_values = {
        "Total Companies": analytics.get("total_companies", 0),
        "Total Contacts": analytics.get("total_contacts", 0),
        "Drafted": analytics.get("drafted", 0),
        "Sent": analytics.get("sent", 0),
        "Replied": analytics.get("replied", 0),
        "Interested": analytics.get("interested", 0),
        "Meetings": analytics.get("meetings", 0),
        "Follow-ups": analytics.get("followups", 0),
        "Not Interested": analytics.get("not_interested", 0),
        "Closed": analytics.get("closed", 0),
        "Response Rate": analytics.get("response_rate", ""),
        "Interest Rate": analytics.get("interest_rate", ""),
    }
    return ExportData(tracker_rows, contact_rows, dashboard_values)


def write_workbook(data: ExportData, config: Config, path: str) -> str:
    """Create/overwrite the tracker workbook. Sync; call via to_thread."""
    wb = Workbook()

    tracker = wb.active
    tracker.title = "Outreach Tracker"
    tracker.append(TRACKER_HEADERS)
    for row in data.tracker_rows:
        tracker.append(row)
    _style_header(tracker, len(TRACKER_HEADERS))
    _size_columns(tracker, [18, 22, 20, 18, 30, 26, 18, 22, 30, 50, 10, 14, 14, 14, 14, 30, 26])
    tracker.freeze_panes = "A2"

    contacts = wb.create_sheet("Contacts")
    contacts.append(CONTACT_HEADERS)
    for row in data.contact_rows:
        contacts.append(row)
    _style_header(contacts, len(CONTACT_HEADERS))
    _size_columns(contacts, [10, 10, 22, 22, 20, 34, 26, 30])
    contacts.freeze_panes = "A2"

    dashboard = wb.create_sheet("Dashboard")
    dashboard.append(["Metric", "Value"])
    for label in DASHBOARD_ROWS:
        dashboard.append([label, data.dashboard_values.get(label, "")])
    _style_header(dashboard, 2)
    _size_columns(dashboard, [22, 14])

    settings = wb.create_sheet("Settings")
    settings.append(["Setting", "Value"])
    for label, value in [
        ("Student Name", config.student_name),
        ("College", config.student_college),
        ("Student Role", config.student_role),
        ("Default Follow-up Days", config.default_followup_days),
        ("Last Updated", now_ts()),
    ]:
        settings.append([label, value])
    _style_header(settings, 2)
    _size_columns(settings, [24, 40])

    os.makedirs(Path(path).parent, exist_ok=True)
    wb.save(path)
    return path


def _style_header(ws, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def _size_columns(ws, widths: list[int]) -> None:
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


class ExcelService:
    def __init__(self, config: Config, repository: Repository) -> None:
        self.config = config
        self.repository = repository

    async def export(self, path: str | None = None) -> str:
        """Pull data from SQLite and (re)build the Excel tracker."""
        target = path or self.config.excel_path
        companies = await self.repository.list_companies()
        contacts = await self.repository.list_contacts()
        outreach = await self.repository.list_outreach()
        analytics = await compute_analytics(self.repository)
        data = build_export_data(companies, contacts, outreach, analytics)
        return await asyncio.to_thread(write_workbook, data, self.config, target)
