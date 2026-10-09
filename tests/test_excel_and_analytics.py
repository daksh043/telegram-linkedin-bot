"""Tests for Excel export and analytics."""

import asyncio

import pytest
from openpyxl import load_workbook

from app.config import Config
from app.database.models import (
    STATUS_DRAFTED,
    STATUS_INTERESTED,
    STATUS_MEETING,
    STATUS_SENT,
)
from app.database.repository import Repository
from app.services.analytics import compute_analytics, rates
from app.services.excel import ExcelService, build_export_data, write_workbook


def make_config(tmp_path) -> Config:
    return Config(
        telegram_bot_token="x",
        ai_api_key="",
        ai_model="gemini-3.8-flash",
        ai_base_url=None,
        database_path=str(tmp_path / "test.db"),
        excel_path=str(tmp_path / "Outreach_Tracker.xlsx"),
        student_name="Daksh Khandagale",
        student_college="Universal AI University, Karjat",
        student_role="Student",
        default_followup_days=7,
        log_level="INFO",
    )


@pytest.fixture
def seeded(tmp_path):
    config = make_config(tmp_path)
    repo = Repository(config.database_path)
    asyncio.run(repo.init())
    cid = asyncio.run(
        repo.add_company(
            name="Acme",
            industry="AI",
            opportunity="Internship",
            personalization_angle="AI hiring push",
        )
    )
    contact_id = asyncio.run(
        repo.add_contact(cid, name="Priya Sharma", role="HR", email="p@acme.com")
    )
    asyncio.run(repo.add_outreach(cid, contact_id, "Email", "Draft msg", STATUS_DRAFTED))
    asyncio.run(repo.add_outreach(cid, contact_id, "LinkedIn", "Sent msg", STATUS_SENT))
    asyncio.run(repo.update_outreach(2, status=STATUS_SENT, date_contacted="2026-10-01")
    )
    asyncio.run(repo.add_outreach(cid, None, "Email", "Met", STATUS_MEETING))
    return config, repo


def test_rates():
    assert rates(0, 0, 0) == ("", "")
    assert rates(10, 5, 2) == ("50%", "20%")


def test_analytics(seeded):
    _, repo = seeded
    analytics = asyncio.run(compute_analytics(repo))
    assert analytics["total_companies"] == 1
    assert analytics["total_contacts"] == 1
    assert analytics["drafted"] == 1
    assert analytics["sent"] == 1
    assert analytics["meetings"] == 1
    # sent=1, replied-or-beyond = meeting(1) → response 1/2, interest 1/2
    assert analytics["response_rate"] == "50%"
    assert analytics["interest_rate"] == "50%"


def test_excel_export(seeded):
    config, repo = seeded
    service = ExcelService(config, repo)
    path = asyncio.run(service.export())

    wb = load_workbook(path)
    assert wb.sheetnames == [
        "Outreach Tracker", "Contacts", "Dashboard", "Settings",
    ]

    tracker = wb["Outreach Tracker"]
    headers = [c.value for c in tracker[1]]
    assert headers[:6] == ["ID", "Company", "Contact Person", "Role", "LinkedIn", "Email"]
    assert "Follow-up Date" in headers and "Status" in headers
    rows = list(tracker.iter_rows(min_row=2, values_only=True))
    assert len(rows) == 3
    by_id = {r[0]: r for r in rows}
    assert by_id[1][1] == "Acme"
    assert by_id[1][2] == "Priya Sharma"
    assert by_id[3][2] in (None, "")  # no contact on the last record

    contacts = wb["Contacts"]
    assert [c.value for c in contacts[1]] == [
        "Contact ID", "Company ID", "Company", "Name", "Role", "LinkedIn",
        "Email", "Notes",
    ]
    assert list(contacts.iter_rows(min_row=2, values_only=True))[0][3] == "Priya Sharma"

    dashboard = dict(
        (r[0], r[1]) for r in wb["Dashboard"].iter_rows(min_row=2, values_only=True)
    )
    assert dashboard["Total Companies"] == 1
    assert dashboard["Sent"] == 1
    assert dashboard["Meetings"] == 1
    assert dashboard["Response Rate"] == "50%"

    settings = dict(
        (r[0], r[1]) for r in wb["Settings"].iter_rows(min_row=2, values_only=True)
    )
    assert settings["Student Name"] == "Daksh Khandagale"
    assert settings["Default Follow-up Days"] == 7
    assert settings["Last Updated"]


def test_build_export_data_handles_missing_relations():
    data = build_export_data(
        companies=[{"id": 1, "name": "Acme", "industry": "AI"}],
        contacts=[],
        outreach=[{
            "id": 9, "company_id": 1, "contact_id": None, "channel": "Email",
            "message": "hi", "status": "Drafted", "date_added": "2026-10-01",
            "date_contacted": None, "follow_up_date": None,
            "response": "", "notes": "",
        }],
        analytics={},
    )
    assert data.tracker_rows[0][1] == "Acme"
    assert data.tracker_rows[0][2] == ""  # no contact → empty, not crash
